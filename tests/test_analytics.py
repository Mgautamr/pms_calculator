from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from engine.analytics import (cagr, max_drawdown, period_returns, sharpe_ratio,
                              summarize, volatility, xirr)

WORKBOOK = Path(__file__).parent / "analytics_reference.xlsx"
wb = load_workbook(WORKBOOK, data_only=True)


def to_date(cell_value):
    return cell_value.date() if hasattr(cell_value, "date") else cell_value


# ---------- load the Excel answer key ----------
sheet_a = wb["A_Series"]
DATES = [to_date(sheet_a.cell(row=r, column=1).value) for r in range(2, 15)]
VALUES = [sheet_a.cell(row=r, column=2).value for r in range(2, 15)]
EXPECTED = {
    "years": sheet_a["H2"].value,
    "cagr": sheet_a["H3"].value,
    "monthly_vol": sheet_a["H4"].value,
    "volatility": sheet_a["H5"].value,
    "max_drawdown": sheet_a["H6"].value,
    "risk_free": sheet_a["H7"].value,
    "sharpe": sheet_a["H8"].value,
}

sheet_b = wb["B_XIRR"]
FLOWS = [(to_date(sheet_b.cell(row=r, column=1).value), sheet_b.cell(row=r, column=2).value)
         for r in range(4, 8)]
TWO_FLOWS = [(to_date(sheet_b.cell(row=r, column=1).value), sheet_b.cell(row=r, column=2).value)
             for r in (15, 16)]
EXPECTED_XIRR = sheet_b["B9"].value
EXPECTED_TWO_FLOW_XIRR = sheet_b["B18"].value


# ---------- tests against Excel ----------
def test_cagr_matches_excel():
    days = (DATES[-1] - DATES[0]).days
    assert cagr(VALUES[0], VALUES[-1], days) == pytest.approx(EXPECTED["cagr"], abs=1e-9)


def test_monthly_and_annual_volatility_match_excel():
    returns = period_returns(VALUES)
    assert volatility(returns, 1) == pytest.approx(EXPECTED["monthly_vol"], abs=1e-9)
    assert volatility(returns, 12) == pytest.approx(EXPECTED["volatility"], abs=1e-9)


def test_max_drawdown_matches_excel():
    assert max_drawdown(VALUES) == pytest.approx(EXPECTED["max_drawdown"], abs=1e-9)


def test_sharpe_matches_excel():
    result = sharpe_ratio(EXPECTED["cagr"], EXPECTED["risk_free"], EXPECTED["volatility"])
    assert result == pytest.approx(EXPECTED["sharpe"], abs=1e-9)


def test_summarize_matches_excel():
    s = summarize(DATES, VALUES, EXPECTED["risk_free"])
    assert s["years"] == pytest.approx(EXPECTED["years"])
    assert s["cagr"] == pytest.approx(EXPECTED["cagr"], abs=1e-9)
    assert s["volatility"] == pytest.approx(EXPECTED["volatility"], abs=1e-9)
    assert s["max_drawdown"] == pytest.approx(EXPECTED["max_drawdown"], abs=1e-9)
    assert s["sharpe"] == pytest.approx(EXPECTED["sharpe"], abs=1e-9)


def test_xirr_matches_excel():
    assert xirr(FLOWS) == pytest.approx(EXPECTED_XIRR, abs=1e-8)


def test_two_flow_xirr_matches_excel_and_equals_cagr():
    assert xirr(TWO_FLOWS) == pytest.approx(EXPECTED_TWO_FLOW_XIRR, abs=1e-8)
    assert xirr(TWO_FLOWS) == pytest.approx(EXPECTED["cagr"], abs=1e-8)


# ---------- edge cases ----------
def test_rising_series_has_zero_drawdown():
    assert max_drawdown([100, 101, 105, 110]) == 0.0


def test_drawdown_uses_the_running_peak_not_the_start():
    # peak 150, trough 90: -40%, even though the series ends above its start
    assert max_drawdown([100, 150, 90, 120]) == pytest.approx(-0.4)


def test_constant_series_has_zero_volatility_and_undefined_sharpe():
    returns = period_returns([100, 100, 100, 100])
    assert volatility(returns, 12) == 0
    with pytest.raises(ValueError):
        sharpe_ratio(0.1, 0.06, 0)


def test_xirr_needs_both_inflows_and_outflows():
    with pytest.raises(ValueError):
        xirr([(date(2025, 1, 1), -100), (date(2026, 1, 1), -50)])


def test_xirr_handles_a_loss():
    rate = xirr([(date(2025, 1, 1), -100), (date(2026, 1, 1), 80)])
    assert rate == pytest.approx(-0.2, abs=1e-8)


def test_xirr_is_order_independent():
    assert xirr(list(reversed(FLOWS))) == pytest.approx(xirr(FLOWS), abs=1e-10)


def test_cagr_rejects_bad_input():
    with pytest.raises(ValueError):
        cagr(0, 100, 365)
    with pytest.raises(ValueError):
        cagr(100, 120, 0)
