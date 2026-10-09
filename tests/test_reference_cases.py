from pathlib import Path
import pytest
from openpyxl import load_workbook

WORKBOOK = Path(__file__).parent / "reference_cases.xlsx"
def load_cases():
    wb = load_workbook(WORKBOOK, data_only=True)
    cases = []
    for name in wb.sheetnames:
        if not (name.startswith("C") and name[1:3].isdigit()):
            continue
        ws = wb[name]
        cases.append({
            "id": name,
            "initial": ws["B4"].value,
            "fixed_fee": ws["B5"].value,
            "perf_fee": ws["B6"].value,
            "hurdle": ws["B7"].value,
            "catch_up": bool(ws["B8"].value),
            "hwm": bool(ws["B9"].value),
            "exit_load": ws["B10"].value,
            "gross_returns": [ws.cell(row=r, column=2).value for r in range(13, 18)],
            "expected_closing": [ws.cell(row=r, column=10).value for r in range(13, 18)],
            "expected_final": ws["B25"].value,
            "expected_total_fees": ws["B28"].value,
        })
    return cases
engine_fees = pytest.importorskip("engine.fees")


@pytest.mark.parametrize("case", load_cases(), ids=lambda c: c["id"])
def test_case(case):
    result = engine_fees.simulate(
        case["initial"],
        case["gross_returns"],
        case["fixed_fee"],
        case["perf_fee"],
        case["hurdle"],
        case["catch_up"],
        case["hwm"],
        case["exit_load"],
    )
    assert result["closing_values"] == pytest.approx(case["expected_closing"], abs=0.01)
    assert result["final_net_value"] == pytest.approx(case["expected_final"], abs=0.01)
    assert result["total_fees"] == pytest.approx(case["expected_total_fees"], abs=0.01)